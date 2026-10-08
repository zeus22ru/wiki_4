#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Адаптер PPTX: рекурсивный обход фигур, таблицы, картинки, диаграммы, заметки."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Dict, List, Optional

from ..base import (
    ExtractionBlock,
    ExtractionResult,
    Locator,
    VisualOccurrence,
    make_occurrence_id,
    make_source_id,
    stable_hash,
)
from ..imaging import detect_mime_from_bytes, probe_image_size

try:  # pragma: no cover
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    PPTX_AVAILABLE = True
except Exception:  # pragma: no cover
    Presentation = None  # type: ignore
    MSO_SHAPE_TYPE = None  # type: ignore
    PPTX_AVAILABLE = False


class PptxSourceAdapter:
    name = "pptx"
    version = "1"
    extensions = (".pptx",)

    def extract(self, source: Any, *, options: Optional[Dict[str, Any]] = None) -> ExtractionResult:
        path = Path(source)
        options = options or {}
        rel_path = str(options.get("relative_path") or path.name)
        source_id = make_source_id(rel_path)
        result = ExtractionResult(
            source_id=source_id, title=path.stem, source_type="pptx",
            source_path=rel_path, adapter_version=self.version,
        )
        if not PPTX_AVAILABLE:
            result.status = "unsupported"
            result.add_diagnostic("pptx_lib_missing", "Не установлен python-pptx")
            return result

        blocks: List[ExtractionBlock] = []
        occurrences: List[VisualOccurrence] = []
        order = 0
        try:
            prs = Presentation(str(path))
            title = ""
            for slide_num, slide in enumerate(prs.slides, 1):
                slide_title = self._slide_title(slide)
                if not title and slide_title:
                    title = slide_title
                for shape in slide.shapes:
                    order = self._handle_shape(
                        shape, slide_num, slide_title, order,
                        blocks, occurrences, result,
                    )
                notes = self._notes_text(slide)
                if notes:
                    blocks.append(ExtractionBlock(
                        block_id=f"b{order}", kind="paragraph", order=order,
                        text=f"[заметки] {notes}", section_path=slide_title,
                        locator=Locator(slide=slide_num, block_order=order,
                                        object_path="notes"),
                        metadata={"context_kind": "notes"},
                    ))
                    order += 1
            if title:
                result.title = title
        except Exception as exc:
            result.status = "failed"
            result.add_diagnostic("pptx_error", str(exc))
            return result

        revision = stable_hash(
            "|".join(b.text for b in blocks),
            "|".join(str(o.occurrence_id) for o in occurrences),
            length=32,
        )
        for occ in occurrences:
            occ.occurrence_id = make_occurrence_id(source_id, revision, occ.locator.key())

        result.blocks = blocks
        result.visual_occurrences = occurrences
        result.source_revision = revision
        result.status = "success" if (blocks or occurrences) else "empty"
        return result

    def _handle_shape(
        self,
        shape: Any,
        slide_num: int,
        slide_title: str,
        order: int,
        blocks: List[ExtractionBlock],
        occurrences: List[VisualOccurrence],
        result: ExtractionResult,
        *,
        depth: int = 0,
    ) -> int:
        if depth > 8:
            return order
        try:
            shape_type = shape.shape_type
        except Exception:
            shape_type = None

        # Группа — рекурсивно.
        if shape_type == MSO_SHAPE_TYPE.GROUP if MSO_SHAPE_TYPE else False:
            for sub in shape.shapes:
                order = self._handle_shape(sub, slide_num, slide_title, order,
                                           blocks, occurrences, result, depth=depth + 1)
            return order

        obj_path = f"slide[{slide_num}]/shape"
        # Таблица.
        if getattr(shape, "has_table", False) and shape.has_table:
            rows = []
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append(ExtractionBlock(
                    block_id=f"b{order}", kind="table", order=order,
                    text="\n".join(rows), section_path=slide_title,
                    locator=Locator(slide=slide_num, block_order=order, object_path=obj_path + "/table"),
                ))
                order += 1
            return order

        # Картинка.
        if shape_type == MSO_SHAPE_TYPE.PICTURE if MSO_SHAPE_TYPE else False:
            data = self._picture_bytes(shape)
            if data:
                w, h = probe_image_size(data)
                occurrences.append(VisualOccurrence(
                    occurrence_id="", source_id=result.source_id,
                    locator=Locator(slide=slide_num, block_order=order,
                                     object_path=obj_path + "/picture"),
                    image_bytes=data, mime_type=detect_mime_from_bytes(data),
                    width=w, height=h, section_path=slide_title,
                    extraction_quality="native",
                ))
                blocks.append(ExtractionBlock(
                    block_id=f"b{order}", kind="visual", order=order,
                    section_path=slide_title,
                    locator=Locator(slide=slide_num, block_order=order),
                ))
                order += 1
            return order

        # Диаграмма.
        if getattr(shape, "has_chart", False) and shape.has_chart:
            chart_text = self._chart_text(shape.chart)
            if chart_text:
                blocks.append(ExtractionBlock(
                    block_id=f"b{order}", kind="paragraph", order=order,
                    text=chart_text, section_path=slide_title,
                    locator=Locator(slide=slide_num, block_order=order, object_path=obj_path + "/chart"),
                    metadata={"native": "chart"},
                ))
                order += 1
            return order

        # Текст.
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                blocks.append(ExtractionBlock(
                    block_id=f"b{order}", kind="paragraph", order=order, text=text,
                    section_path=slide_title,
                    locator=Locator(slide=slide_num, block_order=order, object_path=obj_path),
                ))
                order += 1
        return order

    @staticmethod
    def _picture_bytes(shape: Any) -> Optional[bytes]:
        try:
            image = shape.image
            return image.blob
        except Exception:
            return None

    @staticmethod
    def _chart_text(chart: Any) -> str:
        parts: List[str] = []
        try:
            if chart.has_title:
                parts.append(f"Диаграмма: {chart.chart_title.text_frame.text}")
        except Exception:
            pass
        try:
            for series in chart.plots[0].series:
                cats = list(getattr(series, "categories", []) or [])
                vals = list(getattr(series, "values", []) or [])
                name = getattr(series, "name", "")
                pairs = ", ".join(f"{c}={v}" for c, v in zip(cats, vals))
                parts.append(f"{name}: {pairs}")
        except Exception:
            pass
        return "; ".join(p for p in parts if p)

    @staticmethod
    def _slide_title(slide: Any) -> str:
        try:
            if slide.shapes.title is not None and slide.shapes.title.text.strip():
                return slide.shapes.title.text.strip()
        except Exception:
            pass
        return ""

    @staticmethod
    def _notes_text(slide: Any) -> str:
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                return slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            pass
        return ""
