#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Адаптер XLSX: значения ячеек, формулы, встроенные изображения и диаграммы."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional

from ..base import (
    ExtractionBlock,
    ExtractionResult,
    Locator,
    VisualOccurrence,
    make_occurrence_id,
    make_source_id,
    stable_hash,
)
from ..imaging import detect_mime_from_bytes, probe_image_size

try:  # pragma: no cover
    from openpyxl import load_workbook

    XLSX_AVAILABLE = True
except Exception:  # pragma: no cover
    load_workbook = None  # type: ignore
    XLSX_AVAILABLE = False


class XlsxSourceAdapter:
    name = "xlsx"
    version = "1"
    extensions = (".xlsx", ".xlsm")

    def extract(self, source: Any, *, options: Optional[Dict[str, Any]] = None) -> ExtractionResult:
        path = Path(source)
        options = options or {}
        rel_path = str(options.get("relative_path") or path.name)
        source_id = make_source_id(rel_path)
        result = ExtractionResult(
            source_id=source_id, title=path.stem, source_type="xlsx",
            source_path=rel_path, adapter_version=self.version,
        )
        if not XLSX_AVAILABLE:
            result.status = "unsupported"
            result.add_diagnostic("xlsx_lib_missing", "Не установлен openpyxl")
            return result

        max_rows = int(options.get("max_rows_per_sheet", 20000) or 20000)
        extract_images = bool(options.get("extract_images", True))
        extract_charts = bool(options.get("extract_charts", True))

        blocks: List[ExtractionBlock] = []
        occurrences: List[VisualOccurrence] = []
        order = 0
        try:
            # Значения (кэш) и формулы читаем раздельно (ТЗ §9.4).
            wb_values = load_workbook(path, data_only=True, read_only=False)
            wb_formulas = load_workbook(path, data_only=False, read_only=False)
            try:
                for sheet_name in wb_values.sheetnames:
                    ws = wb_values[sheet_name]
                    ws_f = wb_formulas[sheet_name] if sheet_name in wb_formulas.sheetnames else None
                    rows_text: List[str] = []
                    for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                        if r_idx > max_rows:
                            result.add_diagnostic("row_limit", f"Лист {sheet_name}: > {max_rows} строк")
                            break
                        values = ["" if v is None else str(v) for v in row]
                        if not any(v.strip() for v in values):
                            continue
                        row_text = " | ".join(values)
                        # Формулы отдельно.
                        if ws_f is not None:
                            formulas = []
                            for cell in ws_f[r_idx]:
                                if isinstance(cell.value, str) and cell.value.startswith("="):
                                    cached = ws[cell.coordinate].value
                                    formulas.append(
                                        f"{cell.coordinate}: {cell.value} (значение: {cached if cached is not None else 'нет'})"
                                    )
                            if formulas:
                                row_text += "  ⟨" + "; ".join(formulas) + "⟩"
                        rows_text.append(row_text)
                    if rows_text:
                        blocks.append(ExtractionBlock(
                            block_id=f"b{order}", kind="table", order=order,
                            text=f"Лист: {sheet_name}\n" + "\n".join(rows_text),
                            section_path=sheet_name,
                            locator=Locator(sheet=sheet_name, block_order=order),
                        ))
                        order += 1

                    if extract_images:
                        order = self._extract_images(ws, sheet_name, source_id, order,
                                                     blocks, occurrences, result)
                    if extract_charts:
                        order = self._extract_charts(ws, sheet_name, order, blocks, result)
            finally:
                wb_values.close()
                wb_formulas.close()
        except Exception as exc:
            result.status = "failed"
            result.add_diagnostic("xlsx_error", str(exc))
            return result

        revision = stable_hash(
            "|".join(b.text for b in blocks if b.kind == "table"),
            length=32,
        )
        for occ in occurrences:
            occ.occurrence_id = make_occurrence_id(source_id, revision, occ.locator.key())
        result.blocks = blocks
        result.visual_occurrences = occurrences
        result.source_revision = revision
        result.status = "success" if blocks else "empty"
        return result

    def _extract_images(
        self, ws: Any, sheet_name: str, source_id: str, order: int,
        blocks: List[ExtractionBlock], occurrences: List[VisualOccurrence], result: ExtractionResult,
    ) -> int:
        images = getattr(ws, "_images", None) or []
        for img in images:
            try:
                data = img._data() if callable(getattr(img, "_data", None)) else None
            except Exception:
                data = None
            if not data:
                continue
            anchor = self._anchor_text(getattr(img, "anchor", None))
            w, h = probe_image_size(data)
            occurrences.append(VisualOccurrence(
                occurrence_id="", source_id=source_id,
                locator=Locator(sheet=sheet_name, cell_range=anchor, block_order=order,
                                object_path=f"{sheet_name}/image"),
                image_bytes=data, mime_type=detect_mime_from_bytes(data),
                width=w, height=h, section_path=sheet_name,
                extraction_quality="native",
            ))
            blocks.append(ExtractionBlock(
                block_id=f"b{order}", kind="visual", order=order,
                section_path=sheet_name,
                locator=Locator(sheet=sheet_name, block_order=order, cell_range=anchor),
            ))
            order += 1
        return order

    def _extract_charts(
        self, ws: Any, sheet_name: str, order: int,
        blocks: List[ExtractionBlock], result: ExtractionResult,
    ) -> int:
        charts = getattr(ws, "_charts", None) or []
        for chart in charts:
            text = self._chart_text(chart)
            if text:
                blocks.append(ExtractionBlock(
                    block_id=f"b{order}", kind="paragraph", order=order,
                    text=f"Диаграмма ({sheet_name}): {text}",
                    section_path=sheet_name,
                    locator=Locator(sheet=sheet_name, block_order=order,
                                    object_path=f"{sheet_name}/chart"),
                    metadata={"native": "chart"},
                ))
                order += 1
        return order

    @staticmethod
    def _chart_text(chart: Any) -> str:
        parts: List[str] = []
        try:
            title = getattr(chart, "title", None)
            if title is not None and getattr(title, "tx", None):
                rich = title.tx.rich
                if rich is not None:
                    parts.append(
                        "Заголовок: " + "".join(
                            r.t or "" for p in rich.p for r in (p.r or [])
                        )
                    )
        except Exception:
            pass
        try:
            for series in getattr(chart, "series", []) or []:
                name = ""
                try:
                    name = series.tx.strRef.f if series.tx and series.tx.strRef else ""
                except Exception:
                    name = ""
                ref = ""
                try:
                    if series.val and series.val.numRef:
                        ref = series.val.numRef.f or ""
                except Exception:
                    pass
                if name or ref:
                    parts.append(f"{name}: {ref}".strip(": "))
        except Exception:
            pass
        return "; ".join(p for p in parts if p)

    @staticmethod
    def _anchor_text(anchor: Any) -> Optional[str]:
        try:
            if anchor is None:
                return None
            fr = getattr(anchor, "_from", None)
            if fr is not None:
                return f"R{fr.row + 1}C{fr.col + 1}"
        except Exception:
            pass
        return None
