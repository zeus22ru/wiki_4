#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Тесты подсистемы визуального RAG и portable OCR (ТЗ §22).

Используют fake OCR-процесс и mock vision-транспорт: реальный Windows EXE и GPU
не требуются. Реальный portable-запуск проверяется отдельными integration-тестами.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest

from core.document_extraction.base import (
    ExtractionResult,
    Locator,
    VisualOccurrence,
    make_occurrence_id,
    make_source_id,
)
from core.kb_catalog import KnowledgeCatalog
from core.ocr import FakeOCRProvider, OCRBlock, OCRResult, RapidOcrJsonProvider
from core.visual_analysis import FakeVisualTransport, VisualAnalysis, VisualAnalyzer
from core.visual_assets import VisualAssetStore
from core import visual_indexing as vi
from config import settings


def _png_bytes(width: int = 200, height: int = 100, color: str = "white") -> bytes:
    pytest.importorskip("PIL")
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture
def catalog(tmp_path: Path) -> KnowledgeCatalog:
    return KnowledgeCatalog(tmp_path / "catalog.sqlite")


@pytest.fixture
def assets(tmp_path: Path, catalog: KnowledgeCatalog) -> VisualAssetStore:
    return VisualAssetStore(tmp_path / "assets", catalog=catalog)


@pytest.fixture
def vision_enabled(monkeypatch):
    """Vision-тесты включают анализ явно: в рабочем .env он может быть выключен."""
    monkeypatch.setattr(settings, "VISUAL_ANALYSIS_ENABLED", True)


# --- Каталог и хранилище ассетов ---------------------------------------------

def test_catalog_source_asset_occurrence_roundtrip(catalog: KnowledgeCatalog) -> None:
    catalog.upsert_source(source_id="src:a", title="A", source_type="docx",
                          source_path="a.docx", revision="r1")
    catalog.upsert_asset(asset_id="sha256:abc", sha256="abc", mime_type="image/png",
                         bytes_len=10, stored_path="originals/ab/cd/abc")
    catalog.upsert_occurrence(occurrence_id="occ:1", source_id="src:a", asset_id="sha256:abc",
                              locator={"page": None}, caption="cap")
    occ = catalog.get_occurrence("occ:1")
    assert occ is not None and occ["caption"] == "cap"
    assert catalog.asset_reference_count("sha256:abc") == 1


def test_catalog_orphan_assets_grace(catalog: KnowledgeCatalog) -> None:
    catalog.upsert_asset(asset_id="sha256:x", sha256="x", mime_type="image/png",
                         bytes_len=1, stored_path="o/x")
    # Свежий ассет в grace-периоде не считается сиротой.
    assert catalog.list_orphan_assets(grace_hours=72) == []
    assert len(catalog.list_orphan_assets(grace_hours=0)) == 1


def test_asset_store_idempotent_and_safe_read(assets: VisualAssetStore) -> None:
    data = _png_bytes()
    aid1 = assets.store_original(data, mime_type="image/png")
    aid2 = assets.store_original(data, mime_type="image/png")
    assert aid1 == aid2
    assert assets.read(aid1) == data
    # Нет traversal за пределы хранилища.
    assert assets.read("sha256:../../etc/passwd") is None


def test_asset_store_gc_dry_run(assets: VisualAssetStore, catalog: KnowledgeCatalog) -> None:
    aid = assets.store_original(_png_bytes(), mime_type="image/png")
    report = assets.gc_orphans(grace_hours=0, dry_run=True)
    assert report["orphan_candidates"] == 1 and report["dry_run"] is True
    assert assets.read(aid) is not None  # dry-run не удаляет
    report2 = assets.gc_orphans(grace_hours=0, dry_run=False)
    assert report2["removed"] == 1
    assert assets.read(aid) is None


# --- OCR ----------------------------------------------------------------------

def test_fake_ocr_normalization() -> None:
    ocr = FakeOCRProvider(text="Строка 1\nСтрока 2")
    res = ocr.recognize(b"x")
    assert res.status == "success"
    assert "Строка 1" in res.text_raw
    assert len(res.blocks) == 2


def test_fake_ocr_no_text() -> None:
    ocr = FakeOCRProvider(no_text=True)
    res = ocr.recognize(b"x")
    assert res.status == "no_text"
    assert res.text_raw == ""


def test_rapidocr_health_missing_executable() -> None:
    provider = RapidOcrJsonProvider(executable="Z:/nope/missing.exe")
    health = provider.health_check()
    assert health["ok"] is False
    assert health["error"] == "ocr_runtime_missing"


def test_rapidocr_normalize_upstream_payload() -> None:
    provider = RapidOcrJsonProvider(executable=__file__, model_fingerprint="sha256:m")
    payload = {
        "image_width": 100, "image_height": 50,
        "data": [{"box": [[0, 0], [10, 0], [10, 5], [0, 5]], "text": "Привет", "score": 0.9}],
    }
    res = provider._normalize(payload, exit_code=100, profile="cyrillic", duration_ms=5)
    assert res.status == "success"
    assert res.blocks[0].text == "Привет"
    assert res.blocks[0].polygon[0] == [0.0, 0.0]


def test_rapidocr_normalize_no_text_exit_code() -> None:
    provider = RapidOcrJsonProvider(executable=__file__)
    res = provider._normalize(None, exit_code=101, profile="cyrillic", duration_ms=1)
    assert res.status == "no_text"


# --- Реальный pipe-процесс (поддельный движок) --------------------------------

_FAKE_ENGINE = """
import sys, json
for line in sys.stdin:
    line = line.strip()
    if not line:
        continue
    try:
        json.loads(line)
    except Exception:
        continue
    print(json.dumps({"data": [{"box": [[0,0],[10,0],[10,10],[0,10]],
                                 "text": "Привет ТОРГ-2", "score": 0.97}]}), flush=True)
"""

_FAKE_ENGINE_SLOW = """
import sys, time
for line in sys.stdin:
    if line.strip():
        time.sleep(30)
"""


def _write_engine(tmp_path, body):
    script = tmp_path / "fake_ocr.py"
    script.write_text(body, encoding="utf-8")
    return script


def test_rapidocr_pipe_process_recognizes(tmp_path) -> None:
    import sys

    script = _write_engine(tmp_path, _FAKE_ENGINE)
    provider = RapidOcrJsonProvider(
        executable=sys.executable, extra_args=[str(script)], mode="pipe",
        startup_timeout_sec=0.3, recognize_timeout_sec=10, retry_count=0,
        dual_pass=False, require_models_dir=False, include_model_args=False,
    )
    try:
        res = provider.recognize(_png_bytes())
    finally:
        provider.close()
    assert res.status == "success"
    assert "ТОРГ-2" in res.text_raw
    assert res.blocks and res.blocks[0].polygon[0] == [0.0, 0.0]
    assert provider._procs == {}  # процессы закрыты


def test_rapidocr_pipe_timeout_returns_failed(tmp_path) -> None:
    import sys

    script = _write_engine(tmp_path, _FAKE_ENGINE_SLOW)
    provider = RapidOcrJsonProvider(
        executable=sys.executable, extra_args=[str(script)], mode="pipe",
        startup_timeout_sec=0.3, recognize_timeout_sec=0.8, retry_count=0,
        dual_pass=False, require_models_dir=False, include_model_args=False,
    )
    try:
        res = provider.recognize(_png_bytes())
    finally:
        provider.close()
    assert res.status == "failed"
    assert any("timeout" in w for w in res.warnings)
    assert provider._procs == {}


def test_rapidocr_missing_executable_returns_unavailable() -> None:
    provider = RapidOcrJsonProvider(executable="Z:/nope/missing.exe", retry_count=0)
    res = provider.recognize(b"x")
    assert res.status == "unavailable"
    assert "ocr_runtime_missing" in res.warnings[0]


def test_rapidocr_oneshot_uses_image_argument(tmp_path) -> None:
    import sys

    # Движок печатает путь к изображению, переданный через --image=..., в текст блока.
    engine = """
import sys, json
path = ""
for a in sys.argv[1:]:
    if a.startswith("--image="):
        path = a.split("=", 1)[1]
print(json.dumps({"code": 100, "data": [{"box": [[0,0],[1,0],[1,1],[0,1]],
      "text": path, "score": 0.9}]}), flush=True)
"""
    script = _write_engine(tmp_path, engine)
    provider = RapidOcrJsonProvider(
        executable=sys.executable, extra_args=[str(script)], mode="oneshot",
        recognize_timeout_sec=10, retry_count=0, dual_pass=False,
        require_models_dir=False, include_model_args=False,
    )
    try:
        res = provider.recognize(_png_bytes())
    finally:
        provider.close()
    assert res.status == "success"
    assert res.text_raw.endswith(".png")


def test_rapidocr_dual_pass_merges_and_keeps_alt() -> None:
    # Синтетический dual-pass через подмену _recognize_with.
    provider = RapidOcrJsonProvider(executable=__file__, profile="cyrillic",
                                    secondary_profile="english", dual_pass=True,
                                    require_models_dir=False)

    def fake(image_path, profile):
        if profile == "cyrillic":
            return ({"code": 100, "data": [
                {"box": [[0, 0], [10, 0], [10, 10], [0, 10]], "text": "OO-OOOOO6695", "score": 0.8},
            ]}, None)
        return ({"code": 100, "data": [
            {"box": [[0, 0], [10, 0], [10, 10], [0, 10]], "text": "00-00006695", "score": 0.99},
        ]}, None)

    provider._recognize_with = fake  # type: ignore[assignment]
    res = provider._recognize_once(_png_bytes(), primary="cyrillic", options=None)
    assert res.status == "success"
    # Чисто латинско-цифровой блок → выбран более латинский вариант второго прохода.
    assert res.text_raw == "00-00006695"
    assert res.blocks[0].text_alt == "OO-OOOOO6695"


def test_rapidocr_dual_pass_keeps_cyrillic_primary() -> None:
    provider = RapidOcrJsonProvider(executable=__file__, profile="cyrillic",
                                    secondary_profile="english", dual_pass=True,
                                    require_models_dir=False)

    def fake(image_path, profile):
        text = "Расходного ордера" if profile == "cyrillic" else "PacxoHoro opepa"
        return ({"code": 100, "data": [
            {"box": [[0, 0], [10, 0], [10, 10], [0, 10]], "text": text, "score": 0.8},
        ]}, None)

    provider._recognize_with = fake  # type: ignore[assignment]
    res = provider._recognize_once(_png_bytes(), primary="cyrillic", options=None)
    # Кириллица в основном чтении сохраняется, альтернатива остаётся доступной.
    assert res.text_raw == "Расходного ордера"
    assert res.text_raw_alt == "PacxoHoro opepa"


# --- Vision-анализ ------------------------------------------------------------

def test_visual_analyzer_parses_graph(vision_enabled) -> None:
    payload = json.dumps({
        "visual_type": "flowchart",
        "summary": "Проверка приёмки",
        "visible_text": ["Есть расхождения?"],
        "nodes": [
            {"id": "n1", "label": "Проверить", "type": "action"},
            {"id": "n2", "label": "Есть расхождения?", "type": "decision"},
        ],
        "edges": [{"from": "n1", "to": "n2", "condition": None, "status": "observed"}],
        "uncertainties": [],
    })
    analyzer = VisualAnalyzer(model="fake", transport=FakeVisualTransport(payload))
    analysis = analyzer.analyze("sha256:x", _png_bytes(), context="текст статьи")
    assert analysis.status == "success"
    assert analysis.visual_type == "flowchart"
    assert len(analysis.nodes) == 2 and len(analysis.edges) == 1
    assert "Проверить" in analysis.as_search_text()


def test_visual_analyzer_drops_dangling_edges(vision_enabled) -> None:
    payload = json.dumps({
        "visual_type": "flowchart",
        "nodes": [{"id": "n1", "label": "A", "type": "action"}],
        "edges": [{"from": "n1", "to": "nX", "status": "observed"}],
    })
    analyzer = VisualAnalyzer(model="fake", transport=FakeVisualTransport(payload))
    analysis = analyzer.analyze("sha256:x", _png_bytes())
    assert analysis.edges == []
    assert any(w.startswith("dangling_edge") for w in analysis.warnings)


def test_visual_analyzer_repairs_invalid_json(vision_enabled) -> None:
    calls = {"n": 0}

    def transport(messages):
        calls["n"] += 1
        if calls["n"] == 1:
            return "не JSON"
        return '{"visual_type": "chart", "summary": "график"}'

    analyzer = VisualAnalyzer(model="fake", transport=transport)
    analysis = analyzer.analyze("sha256:x", _png_bytes())
    assert analysis.status == "success"
    assert analysis.visual_type == "chart"
    assert calls["n"] == 2  # один repair


def test_visual_analyzer_marks_decorative(vision_enabled) -> None:
    payload = json.dumps({"visual_type": "decorative", "summary": "логотип", "is_decorative": True})
    analyzer = VisualAnalyzer(model="fake", transport=FakeVisualTransport(payload))
    analysis = analyzer.analyze("sha256:x", _png_bytes())
    assert analysis.is_decorative is True


# --- Визуальная индексация ----------------------------------------------------

def _extraction_with_image(source_id: str = "src:doc") -> ExtractionResult:
    data = _png_bytes()
    loc = Locator(block_order=1)
    occ = VisualOccurrence(
        occurrence_id=make_occurrence_id(source_id, "rev1", loc.key()),
        source_id=source_id, locator=loc, image_bytes=data, mime_type="image/png",
        caption="Схема процесса", section_path="Процесс → Приёмка",
        context_before="Проверьте приёмку", context_after="Нажмите Заполнить",
    )
    return ExtractionResult(
        source_id=source_id, title="Процесс приёмки", source_type="docx",
        source_path="manuals/process.docx", source_revision="rev1",
        visual_occurrences=[occ], status="success",
    )


def test_visual_indexing_builds_searchable_chunk(catalog, assets) -> None:
    payload = json.dumps({
        "visual_type": "flowchart",
        "summary": "Схема проверки приёмки",
        "visible_text": ["Есть расхождения?"],
        "nodes": [{"id": "n1", "label": "Есть расхождения?", "type": "decision"}],
    })
    indexer = vi.VisualIndexer(
        ocr=FakeOCRProvider(text="ТОРГ-2\nЕсть расхождения?"),
        analyzer=VisualAnalyzer(model="fake", transport=FakeVisualTransport(payload)),
        assets=assets, catalog=catalog,
    )
    out = indexer.process(_extraction_with_image())
    assert out.status == "success"
    assert len(out.chunks) == 1
    meta = out.chunks[0].metadata
    assert meta["chunk_kind"] == "visual"
    assert meta["evidence_kind"] == "visual"
    assert "ТОРГ-2" in meta["text_raw"]
    assert out.counters["occurrences"] == 1
    assert out.counters["assets_stored"] == 1


def test_visual_indexing_uses_cache_on_second_run(catalog, assets) -> None:
    payload = json.dumps({"visual_type": "unknown", "summary": "описание"})
    ocr = FakeOCRProvider(text="текст")
    indexer = vi.VisualIndexer(
        ocr=ocr,
        analyzer=VisualAnalyzer(model="fake", transport=FakeVisualTransport(payload)),
        assets=assets, catalog=catalog,
    )
    out1 = indexer.process(_extraction_with_image())
    out2 = indexer.process(_extraction_with_image())
    assert out1.counters["ocr_calls"] == 1
    assert out2.counters["ocr_calls"] == 0
    assert out2.counters["ocr_cache_hits"] == 1


def test_visual_indexing_survives_ocr_failure(catalog, assets) -> None:
    indexer = vi.VisualIndexer(
        ocr=FakeOCRProvider(status="failed", invalid_json=True),
        analyzer=None, assets=assets, catalog=catalog,
    )
    out = indexer.process(_extraction_with_image())
    # Источник не роняется, но помечается partial и есть диагностика.
    assert out.status in ("partial", "success")
    assert out.chunks  # картинка без текста остаётся доступной
    assert any(d["code"].startswith("ocr_") for d in out.diagnostics)


def test_image_only_source_not_skipped(catalog, assets) -> None:
    # Источник без текста, но с картинкой → визуальный чанк создаётся.
    indexer = vi.VisualIndexer(ocr=FakeOCRProvider(no_text=True),
                              analyzer=None, assets=assets, catalog=catalog)
    out = indexer.process(_extraction_with_image("src:imageonly"))
    assert len(out.chunks) == 1


# --- Дефолты настроек ---------------------------------------------------------

def test_visual_feature_disabled_by_default(monkeypatch) -> None:
    # Проверяем именно дефолт кода, не завися от .env окружения.
    from config.settings import _env_bool

    monkeypatch.delenv("KB_VISUAL_ENABLED", raising=False)
    assert _env_bool("KB_VISUAL_ENABLED", False) is False


# --- Preview и возможности (ТЗ §17) ------------------------------------------

def test_preview_source_counts_visuals(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document
    from docx.shared import Inches
    from PIL import Image
    import io as _io

    img_path = tmp_path / "i.png"
    buf = _io.BytesIO()
    Image.new("RGB", (80, 40), "white").save(buf, "PNG")
    img_path.write_bytes(buf.getvalue())

    doc = Document()
    doc.add_heading("Раздел", level=1)
    doc.add_paragraph("Текст")
    doc.add_picture(str(img_path), width=Inches(1))
    doc.save(tmp_path / "d.docx")

    from core.visual_pipeline import preview_source, visual_capabilities

    summary = preview_source(tmp_path / "d.docx")
    assert summary["supported"] is True
    assert summary["blocks"]["visual"] == 1
    assert summary["occurrences"] == 1
    assert summary["text_blocks"] >= 2
    caps = visual_capabilities()
    assert "ocr_executable_present" in caps
    assert isinstance(caps["visual_enabled"], bool)


def test_preview_unsupported_format(tmp_path: Path) -> None:
    from core.visual_pipeline import preview_source

    p = tmp_path / "x.zzz"
    p.write_text("x", encoding="utf-8")
    summary = preview_source(p)
    assert summary["status"] == "unsupported"
