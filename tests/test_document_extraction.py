#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Тесты адаптеров извлечения (ТЗ §22.1)."""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from core.document_extraction import extract_source
from core.document_extraction.adapters import register_default_adapters


@pytest.fixture(autouse=True)
def _register() -> None:
    register_default_adapters()


def _png_bytes() -> bytes:
    pytest.importorskip("PIL")
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (120, 60), "white").save(buf, format="PNG")
    return buf.getvalue()


# --- TXT ----------------------------------------------------------------------

def test_text_adapter_order_and_headings(tmp_path: Path) -> None:
    p = tmp_path / "doc.txt"
    p.write_text("# Заголовок\nПервая строка\nВторая строка\n", encoding="utf-8")
    res = extract_source(p, options={"relative_path": "doc.txt"})
    assert res.status == "success"
    kinds = [b.kind for b in res.blocks]
    assert kinds[0] == "heading"
    assert kinds.count("paragraph") == 2
    assert res.blocks[0].order < res.blocks[1].order


def test_text_adapter_empty(tmp_path: Path) -> None:
    p = tmp_path / "empty.txt"
    p.write_text("   \n\n", encoding="utf-8")
    res = extract_source(p)
    assert res.status == "empty"


# --- HTML ---------------------------------------------------------------------

def test_html_adapter_keeps_dom_order_and_images(tmp_path: Path) -> None:
    (tmp_path / "img.png").write_bytes(_png_bytes())
    html = """
    <html><body><article>
      <h1>Инструкция</h1>
      <p>Шаг 1</p>
      <figure><img src="img.png" alt="Скриншот"><figcaption>Рисунок 1</figcaption></figure>
      <p>Шаг 2</p>
    </article></body></html>
    """
    (tmp_path / "page.html").write_text(html, encoding="utf-8")
    res = extract_source(tmp_path / "page.html", options={"relative_path": "page.html"})
    assert res.status == "success"
    assert len(res.visual_occurrences) == 1
    occ = res.visual_occurrences[0]
    assert occ.caption == "Рисунок 1"
    assert occ.alt == "Скриншот"
    assert occ.image_bytes and len(occ.image_bytes) > 0
    # Порядок: heading, Шаг 1, visual, Шаг 2
    orders = [(b.kind, b.text) for b in res.blocks]
    assert orders[0][0] == "heading"
    assert any(k == "visual" for k, _ in orders)


def test_html_adapter_inline_data_image(tmp_path: Path) -> None:
    import base64

    b64 = base64.b64encode(_png_bytes()).decode()
    html = f'<article><p>текст</p><img src="data:image/png;base64,{b64}"></article>'
    (tmp_path / "d.html").write_text(html, encoding="utf-8")
    res = extract_source(tmp_path / "d.html")
    assert len(res.visual_occurrences) == 1
    assert res.visual_occurrences[0].image_bytes


def test_html_adapter_strikethrough_marks_stale(tmp_path: Path) -> None:
    html = "<article><p>Актуально</p><p><del>Устаревший шаг</del></p></article>"
    (tmp_path / "s.html").write_text(html, encoding="utf-8")
    res = extract_source(tmp_path / "s.html")
    texts = " ".join(b.text for b in res.blocks)
    assert "[УСТАРЕЛО:" in texts


def test_html_adapter_missing_local_asset_partial(tmp_path: Path) -> None:
    html = '<article><p>текст</p><img src="missing.png"></article>'
    (tmp_path / "m.html").write_text(html, encoding="utf-8")
    res = extract_source(tmp_path / "m.html")
    assert res.status == "partial"
    assert any(d["code"] == "asset_missing" for d in res.diagnostics)


# --- DOCX ---------------------------------------------------------------------

def _make_docx(tmp_path: Path) -> Path:
    docx = pytest.importorskip("docx")
    from docx.shared import Inches

    doc = docx.Document()
    doc.add_heading("Раздел 1", level=1)
    doc.add_paragraph("Первый абзац")
    doc.add_picture(str(tmp_path / "img.png"), width=Inches(1.5))
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    doc.add_paragraph("Абзац после таблицы")
    path = tmp_path / "doc.docx"
    doc.save(path)
    return path


def test_docx_adapter_order_table_image(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    (tmp_path / "img.png").write_bytes(_png_bytes())
    path = _make_docx(tmp_path)
    res = extract_source(path, options={"relative_path": "doc.docx"})
    assert res.status == "success"
    kinds = [b.kind for b in res.blocks]
    assert "heading" in kinds and "table" in kinds and "visual" in kinds
    # Таблица идёт после изображения в порядке документа.
    assert kinds.index("visual") < kinds.index("table")
    assert len(res.visual_occurrences) == 1
    assert res.title == "Раздел 1" or res.title == "doc"


def test_docx_adapter_heading_section_path(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    (tmp_path / "img.png").write_bytes(_png_bytes())
    path = _make_docx(tmp_path)
    res = extract_source(path)
    table_block = next(b for b in res.blocks if b.kind == "table")
    assert "Раздел 1" in table_block.section_path


# --- Изображения --------------------------------------------------------------

def test_image_adapter_standalone(tmp_path: Path) -> None:
    p = tmp_path / "shot.png"
    p.write_bytes(_png_bytes())
    res = extract_source(p, options={"relative_path": "shot.png"})
    assert res.status == "success"
    assert len(res.visual_occurrences) == 1
    occ = res.visual_occurrences[0]
    assert occ.width == 120 and occ.height == 60
    assert occ.mime_type == "image/png"


def test_image_adapter_svg(tmp_path: Path) -> None:
    svg = b'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"><rect/></svg>'
    p = tmp_path / "d.svg"
    p.write_bytes(svg)
    res = extract_source(p)
    assert res.visual_occurrences[0].mime_type == "image/svg+xml"


# --- Неподдерживаемый формат --------------------------------------------------

def test_unsupported_format_explicit(tmp_path: Path) -> None:
    from core.document_extraction.base import UnsupportedFormat

    p = tmp_path / "file.zzz"
    p.write_text("x", encoding="utf-8")
    with pytest.raises(UnsupportedFormat):
        extract_source(p)
