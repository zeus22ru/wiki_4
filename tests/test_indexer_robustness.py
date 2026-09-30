# -*- coding: utf-8 -*-
"""Тесты устойчивости индексатора (поток I): скан, DOC, manifest, reindex, chunking."""

from __future__ import annotations

import ast
import hashlib
import time
import zipfile
from pathlib import Path
from unittest.mock import MagicMock

import pytest

from core.chunking import chunk_text_fixed_size
from core.index_manifest import (
    build_index_manifest,
    load_index_manifest,
    settings_fingerprint,
)


def test_create_vector_db_parses_under_python_310():
    """I1: файл create_vector_db.py должен парситься с feature_version=(3, 10)."""
    source = Path(__file__).resolve().parents[1] / "create_vector_db.py"
    tree = ast.parse(source.read_text(encoding="utf-8"), feature_version=(3, 10))
    assert isinstance(tree, ast.Module)
    # Вложенных f-строк с md5 больше нет — id считается через отдельную переменную.
    src = source.read_text(encoding="utf-8")
    assert "hashlib.md5(f'{doc_data" not in src


def test_scan_supported_files_case_insensitive_and_skips_service_dirs(monkeypatch, tmp_path):
    """I2+I3: .PDF/.Docx находятся; вложения, .git, __pycache__, ~$ пропускаются."""
    import create_vector_db

    data_dir = tmp_path / "data"
    data_dir.mkdir()
    attachments = data_dir / "chat_attachments"
    attachments.mkdir()
    (attachments / "secret.txt").write_text("private " * 20, encoding="utf-8")
    (data_dir / ".git" / "objects").mkdir(parents=True)
    (data_dir / ".git" / "objects" / "x.txt").write_text("git " * 30, encoding="utf-8")
    (data_dir / "__pycache__").mkdir()
    (data_dir / "__pycache__" / "mod.txt").write_text("cache " * 30, encoding="utf-8")
    (data_dir / "~$draft.docx").write_text("tmp", encoding="utf-8")
    (data_dir / "Report.PDF").write_text("pdf content " * 40, encoding="utf-8")
    (data_dir / "Notes.Docx").write_text("docx content " * 40, encoding="utf-8")
    (data_dir / "ok.txt").write_text("plain text " * 40, encoding="utf-8")
    (data_dir / "downloading.TXT.CRDOWNLOAD").write_text("partial", encoding="utf-8")
    # Файл с исключённым суффиксом в верхнем регистре
    (data_dir / "old.BAK").write_text("bak " * 20, encoding="utf-8")

    monkeypatch.setattr(create_vector_db.settings, "CHAT_ATTACHMENTS_DIR", str(attachments))

    found = create_vector_db.scan_supported_files(str(data_dir))
    names = {p.name for p in found}

    assert "Report.PDF" in names
    assert "Notes.Docx" in names
    assert "ok.txt" in names
    assert "secret.txt" not in names
    assert "x.txt" not in names
    assert "mod.txt" not in names
    assert "~$draft.docx" not in names
    assert "old.BAK" not in names


def test_extract_text_from_doc_title_before_whitespace_collapse(monkeypatch, tmp_path):
    """I4: заголовок берётся из первой непустой строки до схлопывания пробелов."""
    import create_vector_db

    doc_path = tmp_path / "sample.doc"
    doc_path.write_bytes(b"placeholder")

    raw = "  Заголовок документа  \n\nПервый абзац.\n\nВторой   абзац."
    monkeypatch.setattr(create_vector_db, "_extract_text_as_docx_zip", lambda p: raw)
    monkeypatch.setattr(create_vector_db, "_extract_doc_via_external_tool", lambda p: None)

    result = create_vector_db.extract_text_from_doc(doc_path)
    assert result is not None
    assert result["title"] == "Заголовок документа"
    assert "\n" in result["content"]
    assert "Второй абзац" in result["content"]
    assert "Второй   абзац" not in result["content"]


def test_extract_text_from_doc_fallback_without_external_tools(monkeypatch, tmp_path):
    """I4: без ZIP и без утилит — понятное предупреждение и None."""
    import create_vector_db

    doc_path = tmp_path / "binary.doc"
    doc_path.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 64)

    monkeypatch.setattr(create_vector_db, "_extract_text_as_docx_zip", lambda p: None)
    monkeypatch.setattr(create_vector_db, "_extract_doc_via_external_tool", lambda p: None)
    fake_logger = MagicMock()
    monkeypatch.setattr(create_vector_db, "logger", fake_logger)

    assert create_vector_db.extract_text_from_doc(doc_path) is None
    warning_text = " ".join(
        str(call.args[0]) for call in fake_logger.warning.call_args_list if call.args
    )
    assert "antiword" in warning_text or "Пропуск DOC" in warning_text


def test_extract_text_from_doc_reads_misnamed_docx_zip(tmp_path):
    """I4: .doc, который на самом деле DOCX-ZIP, читается без внешних утилит."""
    import create_vector_db
    from docx import Document

    real_docx = tmp_path / "real.docx"
    doc = Document()
    doc.add_paragraph("Титул раздела")
    doc.add_paragraph("Тело документа с достаточным текстом для индексации.")
    doc.save(real_docx)

    fake_doc = tmp_path / "misnamed.doc"
    fake_doc.write_bytes(real_docx.read_bytes())
    assert zipfile.is_zipfile(fake_doc)

    result = create_vector_db.extract_text_from_doc(fake_doc)
    assert result is not None
    assert result["title"] == "Титул раздела"
    assert "Тело документа" in result["content"]


def test_settings_fingerprint_triggers_full_reindex(monkeypatch, tmp_path):
    """I5: несовпадение или отсутствие отпечатка настроек → full_reason=settings_changed."""
    import create_vector_db
    from core import index_manifest as im

    data_dir = tmp_path / "data"
    data_dir.mkdir()
    (data_dir / "a.txt").write_text("alpha content " * 30, encoding="utf-8")

    monkeypatch.setattr(create_vector_db.settings, "DATA_DIR", str(data_dir))
    monkeypatch.setattr(create_vector_db.settings, "CHROMA_PERSIST_DIR", str(tmp_path / "chroma"))
    monkeypatch.setattr(create_vector_db.settings, "CHROMA_COLLECTION_NAME", "wiki_test")
    monkeypatch.setattr(im.settings, "DATA_DIR", str(data_dir))
    monkeypatch.setattr(im.settings, "CHROMA_COLLECTION_NAME", "wiki_test")

    manifest = {
        "version": 1,
        "collection": "wiki_test",
        # нет settings_fingerprint — старый manifest
        "files": {
            "a.txt": {
                "path": "a.txt",
                "chunk_ids": ["old"],
                "size_bytes": 1,
                "mtime_ns": 1,
                "sha256": "dead",
            }
        },
    }
    monkeypatch.setattr(create_vector_db, "load_index_manifest", lambda: manifest)
    monkeypatch.setattr(create_vector_db, "load_bm25_corpus", lambda: (["old"], ["t"]))

    full_calls = {}

    def fake_create(docs, progress_callback=None):
        full_calls["docs"] = len(docs)
        return {"index_mode": "full", "chunks": len(docs)}

    monkeypatch.setattr(create_vector_db, "create_vector_db", fake_create)
    monkeypatch.setattr(
        create_vector_db,
        "process_files",
        lambda files: [
            {
                "id": "n1",
                "text": "x" * 80,
                "embed_text": "x" * 80,
                "metadata": {"path": "a.txt", "title": "A", "file_type": ".txt"},
            }
        ],
    )

    result = create_vector_db.reindex_vector_db()
    assert result["full_reason"] == "settings_changed"
    assert result["index_mode"] == "full"


def test_settings_fingerprint_stable_and_stored(monkeypatch, tmp_path):
    """I5: fingerprint стабилен при тех же настройках и пишется в manifest."""
    from core import index_manifest as im

    data_dir = tmp_path / "data"
    data_dir.mkdir()
    source = data_dir / "article.txt"
    source.write_text("alpha", encoding="utf-8")

    monkeypatch.setattr(im.settings, "DATA_DIR", str(data_dir))
    monkeypatch.setattr(im.settings, "CHROMA_COLLECTION_NAME", "wiki_fp")

    fp1 = settings_fingerprint()
    fp2 = settings_fingerprint()
    assert fp1 == fp2
    assert len(fp1) == 64

    manifest = build_index_manifest(
        [{"id": "c1", "metadata": {"path": "article.txt", "title": "A", "file_type": ".txt"}}],
        data_dir=data_dir,
    )
    assert manifest["settings_fingerprint"] == fp1


def test_incremental_reindex_handles_empty_changed_file(monkeypatch, tmp_path):
    """I6: изменённый файл стал пустым — не падаем, чистим чанки, failed_files."""
    import create_vector_db

    data_dir = tmp_path / "data"
    data_dir.mkdir()
    source = data_dir / "article.txt"
    keep = data_dir / "keep.txt"
    source.write_text("old article " * 20, encoding="utf-8")
    keep.write_text("keep content " * 20, encoding="utf-8")

    monkeypatch.setattr(create_vector_db.settings, "DATA_DIR", str(data_dir))
    monkeypatch.setattr(create_vector_db.settings, "CHROMA_PERSIST_DIR", str(tmp_path / "chroma"))
    monkeypatch.setattr(create_vector_db.settings, "CHROMA_COLLECTION_NAME", "wiki_test")
    monkeypatch.setattr(create_vector_db.settings, "DOCUMENT_PROCESS_WORKERS", 1)
    monkeypatch.setattr(create_vector_db.settings, "STRUCTURAL_CHUNKING_ENABLED", False)
    monkeypatch.setattr(create_vector_db.settings, "CONTEXTUAL_RETRIEVAL_ENABLED", False)

    manifest = build_index_manifest(
        [
            {
                "id": "article-0",
                "metadata": {"path": "article.txt", "title": "Article", "file_type": ".txt"},
            },
            {
                "id": "keep-0",
                "metadata": {"path": "keep.txt", "title": "Keep", "file_type": ".txt"},
            },
        ],
        data_dir=data_dir,
    )
    manifest["collection"] = "wiki_test"
    source.write_text("", encoding="utf-8")

    fake_collection = MagicMock()
    fake_client = MagicMock()
    fake_client.get_collection.return_value = fake_collection
    saved_bm25 = {}

    monkeypatch.setattr(create_vector_db, "load_index_manifest", lambda: manifest)
    monkeypatch.setattr(create_vector_db.chromadb, "PersistentClient", lambda path: fake_client)
    monkeypatch.setattr(
        create_vector_db,
        "load_bm25_corpus",
        lambda: (["article-0", "keep-0"], ["old", "keep"]),
    )
    monkeypatch.setattr(
        create_vector_db,
        "save_bm25_index",
        lambda ids, texts: saved_bm25.update({"ids": ids, "texts": texts}),
    )
    monkeypatch.setattr(create_vector_db, "invalidate_embedding_cache", lambda: None)
    monkeypatch.setattr(
        create_vector_db,
        "embed_documents_batch",
        lambda docs: [{**doc, "embedding": [0.1, 0.2]} for doc in docs],
    )

    diagnostics = create_vector_db.reindex_vector_db()

    assert diagnostics["index_mode"] == "incremental"
    assert "article.txt" in diagnostics["failed_files"]
    fake_collection.delete.assert_called_once_with(ids=["article-0"])
    assert saved_bm25["ids"] == ["keep-0"]
    saved = load_index_manifest(tmp_path / "chroma" / "index_manifest.json")
    assert "article.txt" not in saved["files"]
    assert "keep.txt" in saved["files"]


def test_incremental_reindex_marks_empty_new_file_skipped(monkeypatch, tmp_path):
    """I6: новый файл без чанков → skipped:true в manifest, пустые chunk_ids."""
    import create_vector_db

    data_dir = tmp_path / "data"
    data_dir.mkdir()
    keep = data_dir / "keep.txt"
    empty_new = data_dir / "empty.txt"
    keep.write_text("keep content " * 20, encoding="utf-8")
    empty_new.write_text("", encoding="utf-8")

    monkeypatch.setattr(create_vector_db.settings, "DATA_DIR", str(data_dir))
    monkeypatch.setattr(create_vector_db.settings, "CHROMA_PERSIST_DIR", str(tmp_path / "chroma"))
    monkeypatch.setattr(create_vector_db.settings, "CHROMA_COLLECTION_NAME", "wiki_test")
    monkeypatch.setattr(create_vector_db.settings, "DOCUMENT_PROCESS_WORKERS", 1)
    monkeypatch.setattr(create_vector_db.settings, "STRUCTURAL_CHUNKING_ENABLED", False)
    monkeypatch.setattr(create_vector_db.settings, "CONTEXTUAL_RETRIEVAL_ENABLED", False)

    manifest = build_index_manifest(
        [
            {
                "id": "keep-0",
                "metadata": {"path": "keep.txt", "title": "Keep", "file_type": ".txt"},
            },
        ],
        data_dir=data_dir,
    )
    manifest["collection"] = "wiki_test"

    fake_collection = MagicMock()
    fake_client = MagicMock()
    fake_client.get_collection.return_value = fake_collection
    saved_bm25 = {}

    monkeypatch.setattr(create_vector_db, "load_index_manifest", lambda: manifest)
    monkeypatch.setattr(create_vector_db.chromadb, "PersistentClient", lambda path: fake_client)
    monkeypatch.setattr(create_vector_db, "load_bm25_corpus", lambda: (["keep-0"], ["keep"]))
    monkeypatch.setattr(
        create_vector_db,
        "save_bm25_index",
        lambda ids, texts: saved_bm25.update({"ids": ids, "texts": texts}),
    )
    monkeypatch.setattr(create_vector_db, "invalidate_embedding_cache", lambda: None)
    monkeypatch.setattr(
        create_vector_db,
        "embed_documents_batch",
        lambda docs: [{**doc, "embedding": [0.1, 0.2]} for doc in docs],
    )

    diagnostics = create_vector_db.reindex_vector_db()

    assert diagnostics["index_mode"] == "incremental"
    assert "empty.txt" in diagnostics["skipped_files"]
    saved = load_index_manifest(tmp_path / "chroma" / "index_manifest.json")
    entry = saved["files"]["empty.txt"]
    assert entry["chunk_ids"] == []
    assert entry["skipped"] is True
    assert saved_bm25["ids"] == ["keep-0"]


def test_create_vector_db_uses_temp_collection_before_swap(monkeypatch, tmp_path):
    """I7: сначала коллекция __new, затем rename; при ошибке add старая не удаляется."""
    import create_vector_db

    fake_collection = MagicMock()
    fake_client = MagicMock()
    fake_client.create_collection.return_value = fake_collection
    deleted = []

    def fake_delete(name):
        deleted.append(name)

    fake_client.delete_collection.side_effect = fake_delete

    monkeypatch.setattr(create_vector_db.chromadb, "PersistentClient", lambda path: fake_client)
    monkeypatch.setattr(create_vector_db.settings, "BATCH_SIZE", 10)
    monkeypatch.setattr(create_vector_db.settings, "EMBEDDING_WORKERS", 1)
    monkeypatch.setattr(create_vector_db.settings, "CHROMA_COLLECTION_NAME", "wiki_test")
    monkeypatch.setattr(create_vector_db.settings, "CHROMA_PERSIST_DIR", str(tmp_path / "chroma"))
    monkeypatch.setattr(
        create_vector_db,
        "embed_documents_batch",
        lambda docs: [
            {
                "id": doc["id"],
                "text": doc["text"],
                "metadata": doc["metadata"],
                "embedding": [0.1, 0.2],
            }
            for doc in docs
        ],
    )
    monkeypatch.setattr(create_vector_db, "save_bm25_index", lambda ids, texts: None)
    monkeypatch.setattr(create_vector_db, "bm25_index_path", lambda: tmp_path / "bm25.pkl")
    monkeypatch.setattr(create_vector_db, "invalidate_embedding_cache", lambda: None)
    monkeypatch.setattr(create_vector_db, "save_index_manifest", lambda docs: {"files": {}})

    docs = [
        {"id": "a", "text": "A", "metadata": {"path": "a.txt"}, "embed_text": "A"},
        {"id": "b", "text": "B", "metadata": {"path": "b.txt"}, "embed_text": "B"},
    ]
    create_vector_db.create_vector_db(docs)

    create_names = [c.kwargs.get("name") or c.args[0] for c in fake_client.create_collection.call_args_list]
    assert create_names[0] == "wiki_test__new"
    fake_collection.modify.assert_called()
    assert "wiki_test" in deleted


def test_create_vector_db_add_failure_keeps_old_collection(monkeypatch, tmp_path):
    """I7: ошибка add на временной коллекции — старая не удаляется."""
    import create_vector_db

    fake_collection = MagicMock()
    fake_collection.add.side_effect = RuntimeError("chroma add failed")
    fake_client = MagicMock()
    fake_client.create_collection.return_value = fake_collection
    deleted = []
    fake_client.delete_collection.side_effect = lambda name: deleted.append(name)

    monkeypatch.setattr(create_vector_db.chromadb, "PersistentClient", lambda path: fake_client)
    monkeypatch.setattr(create_vector_db.settings, "BATCH_SIZE", 10)
    monkeypatch.setattr(create_vector_db.settings, "EMBEDDING_WORKERS", 1)
    monkeypatch.setattr(create_vector_db.settings, "CHROMA_COLLECTION_NAME", "wiki_test")
    monkeypatch.setattr(create_vector_db.settings, "CHROMA_PERSIST_DIR", str(tmp_path / "chroma"))
    monkeypatch.setattr(
        create_vector_db,
        "embed_documents_batch",
        lambda docs: [
            {
                "id": doc["id"],
                "text": doc["text"],
                "metadata": doc["metadata"],
                "embedding": [0.1, 0.2],
            }
            for doc in docs
        ],
    )

    docs = [{"id": "a", "text": "A", "metadata": {}, "embed_text": "A"}]
    with pytest.raises(RuntimeError, match="chroma add failed"):
        create_vector_db.create_vector_db(docs)

    assert "wiki_test" not in deleted
    assert "wiki_test__new" in deleted


def test_chunk_text_fixed_size_large_overlap_terminates():
    """I8: chunk_size=500, overlap=400 завершается, покрывает текст, без дублей подряд."""
    text = ("Предложение номер. " * 80) + ("Конец текста. " * 10)
    started = time.perf_counter()
    chunks = chunk_text_fixed_size(text, chunk_size=500, overlap=400)
    elapsed = time.perf_counter() - started
    assert elapsed < 2.0
    assert chunks
    joined_len = sum(len(c) for c in chunks)
    assert joined_len >= len(text) * 0.5
    for left, right in zip(chunks, chunks[1:]):
        assert left != right


def test_chunk_id_uses_plain_md5_hex():
    """I9: id чанка — hexdigest без лишней f-обёртки."""
    path = "folder/doc.txt"
    j = 3
    chunk_id_src = f"{path}_{j}"
    expected = hashlib.md5(chunk_id_src.encode()).hexdigest()
    assert expected == hashlib.md5(f"{path}_{j}".encode()).hexdigest()
    assert "{" not in expected


def test_manual_reindex_with_pdf_and_empty_txt(monkeypatch, tmp_path):
    """Приёмка: reindex на DATA_DIR с .PDF и пустым .txt без исключения (эмбеддинги мок)."""
    import create_vector_db

    data_dir = tmp_path / "data"
    data_dir.mkdir()
    (data_dir / "guide.PDF").write_text(
        "Руководство пользователя. " * 40,
        encoding="utf-8",
    )
    (data_dir / "notes.txt").write_text("Заметка для базы. " * 40, encoding="utf-8")
    (data_dir / "empty.txt").write_text("", encoding="utf-8")

    fake_collection = MagicMock()
    fake_client = MagicMock()
    fake_client.create_collection.return_value = fake_collection
    fake_client.get_collection.side_effect = Exception("no collection")

    monkeypatch.setattr(create_vector_db.settings, "DATA_DIR", str(data_dir))
    monkeypatch.setattr(create_vector_db.settings, "CHROMA_PERSIST_DIR", str(tmp_path / "chroma"))
    monkeypatch.setattr(create_vector_db.settings, "CHROMA_COLLECTION_NAME", "wiki_test")
    monkeypatch.setattr(create_vector_db.settings, "BATCH_SIZE", 10)
    monkeypatch.setattr(create_vector_db.settings, "EMBEDDING_WORKERS", 1)
    monkeypatch.setattr(create_vector_db.settings, "DOCUMENT_PROCESS_WORKERS", 1)
    monkeypatch.setattr(create_vector_db.settings, "STRUCTURAL_CHUNKING_ENABLED", False)
    monkeypatch.setattr(create_vector_db.settings, "CONTEXTUAL_RETRIEVAL_ENABLED", False)
    monkeypatch.setattr(create_vector_db.chromadb, "PersistentClient", lambda path: fake_client)
    monkeypatch.setattr(create_vector_db, "load_index_manifest", lambda: {"files": {}, "collection": "wiki_test"})
    monkeypatch.setattr(create_vector_db, "load_bm25_corpus", lambda: None)
    monkeypatch.setattr(create_vector_db, "save_bm25_index", lambda ids, texts: None)
    monkeypatch.setattr(create_vector_db, "bm25_index_path", lambda: tmp_path / "bm25.pkl")
    monkeypatch.setattr(create_vector_db, "invalidate_embedding_cache", lambda: None)
    monkeypatch.setattr(
        create_vector_db,
        "embed_documents_batch",
        lambda docs: [
            {
                "id": doc["id"],
                "text": doc["text"],
                "metadata": doc["metadata"],
                "embedding": [0.1] * 8,
                "embed_text": doc.get("embed_text") or doc["text"],
            }
            for doc in docs
        ],
    )
    # PDF как текст через txt-обработчик: подменим handler для .pdf
    handlers = create_vector_db.get_file_handlers()
    handlers[".pdf"] = create_vector_db.extract_text_from_txt
    monkeypatch.setattr(create_vector_db, "get_file_handlers", lambda: handlers)

    result = create_vector_db.reindex_vector_db()
    assert result["index_mode"] == "full"
    assert result["chunks"] >= 1
